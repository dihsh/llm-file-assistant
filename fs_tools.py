"""
fs_tools.py - Core File System Tools for LLM Function Calling

Provides structured, robust file system operations tailored for reading,
listing, writing, and searching documents (PDF, DOCX, TXT) with full metadata
and error resilience.
"""

import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

# Document parsers
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None


def read_file(filepath: str) -> Dict[str, Any]:
    """
    Read document files (.pdf, .txt, .docx) and extract text content.

    Args:
        filepath: Absolute or relative path to the file.

    Returns:
        A structured dictionary containing:
        - status: "success" or "error"
        - filepath: Normalized string path
        - content: Extracted plain text content
        - metadata: Dictionary of file properties (file_name, extension, size_bytes, modified_date, page_count)
        - error: Error description if status is "error", else None
    """
    path = Path(filepath).resolve()

    if not path.exists():
        return {
            "status": "error",
            "filepath": str(path),
            "content": "",
            "metadata": {},
            "error": f"File not found: {path}",
        }

    if not path.is_file():
        return {
            "status": "error",
            "filepath": str(path),
            "content": "",
            "metadata": {},
            "error": f"Path is not a regular file: {path}",
        }

    stat = path.stat()
    ext = path.suffix.lower()
    metadata: Dict[str, Any] = {
        "file_name": path.name,
        "extension": ext,
        "size_bytes": stat.st_size,
        "modified_date": datetime.fromtimestamp(stat.st_mtime).isoformat(),
        "page_count": None,
    }

    try:
        content = ""
        if ext == ".txt":
            # Attempt UTF-8 first, fallback to latin-1
            try:
                content = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                content = path.read_text(encoding="latin-1")

        elif ext == ".pdf":
            if PdfReader is None:
                raise ImportError("pypdf package is required to read PDF files.")
            reader = PdfReader(str(path))
            metadata["page_count"] = len(reader.pages)
            text_parts = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
            content = "\n\n".join(text_parts)

        elif ext == ".docx":
            if docx is None:
                raise ImportError("python-docx package is required to read DOCX files.")
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract text inside tables if any
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            content = "\n".join(paragraphs)

        else:
            return {
                "status": "error",
                "filepath": str(path),
                "content": "",
                "metadata": metadata,
                "error": f"Unsupported file extension '{ext}'. Supported: .pdf, .docx, .txt",
            }

        return {
            "status": "success",
            "filepath": str(path),
            "content": content.strip(),
            "metadata": metadata,
            "error": None,
        }

    except Exception as exc:
        return {
            "status": "error",
            "filepath": str(path),
            "content": "",
            "metadata": metadata,
            "error": f"Failed reading {path.name}: {str(exc)}",
        }


def list_files(directory: str, extension: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    List all files in a directory with optional extension filtering.

    Args:
        directory: Directory path to scan.
        extension: Optional file extension filter (e.g., '.pdf', 'txt', '.docx').

    Returns:
        List of structured file metadata objects sorted alphabetically by name.
    """
    dir_path = Path(directory).resolve()

    if not dir_path.exists() or not dir_path.is_dir():
        return []

    # Normalize extension (e.g. 'pdf' -> '.pdf', '.PDF' -> '.pdf')
    target_ext = None
    if extension:
        target_ext = extension.strip().lower()
        if not target_ext.startswith("."):
            target_ext = f".{target_ext}"

    files_info = []
    try:
        for entry in dir_path.iterdir():
            if entry.is_file():
                entry_ext = entry.suffix.lower()
                if target_ext is None or entry_ext == target_ext:
                    stat = entry.stat()
                    files_info.append({
                        "name": entry.name,
                        "path": str(entry.resolve()),
                        "size_bytes": stat.st_size,
                        "modified_date": datetime.fromtimestamp(stat.st_mtime).isoformat(),
                        "extension": entry_ext,
                    })
    except Exception:
        return []

    return sorted(files_info, key=lambda x: x["name"].lower())


def write_file(filepath: str, content: str) -> Dict[str, Any]:
    """
    Write text content to a file, automatically creating parent directories.

    Args:
        filepath: Destination file path.
        content: Text content to write.

    Returns:
        Structured response with status, filepath, bytes written, and message.
    """
    path = Path(filepath).resolve()

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        encoded_content = content.encode("utf-8")
        path.write_bytes(encoded_content)

        return {
            "status": "success",
            "filepath": str(path),
            "bytes_written": len(encoded_content),
            "message": f"Successfully wrote {len(encoded_content)} bytes to {path.name}",
        }
    except Exception as exc:
        return {
            "status": "error",
            "filepath": str(path),
            "bytes_written": 0,
            "message": f"Error writing to {filepath}: {str(exc)}",
        }


def search_in_file(filepath: str, keyword: str) -> Dict[str, Any]:
    """
    Search for a keyword inside a document (.pdf, .docx, .txt), returning matches
    with contextual surrounding snippets. Case-insensitive.

    Args:
        filepath: Path to the target document.
        keyword: Search query string.

    Returns:
        Structured dictionary with match count and list of contextual snippets.
    """
    read_result = read_file(filepath)
    if read_result["status"] != "success":
        return {
            "status": "error",
            "filepath": filepath,
            "keyword": keyword,
            "total_matches": 0,
            "matches": [],
            "error": read_result.get("error", "Unable to read file"),
        }

    content = read_result["content"]
    if not keyword or not keyword.strip():
        return {
            "status": "error",
            "filepath": filepath,
            "keyword": keyword,
            "total_matches": 0,
            "matches": [],
            "error": "Keyword must not be empty.",
        }

    # Case-insensitive regex search
    pattern = re.compile(re.escape(keyword), re.IGNORECASE)
    matches = []

    lines = content.splitlines()
    line_offsets = []
    current_offset = 0
    for line in lines:
        line_offsets.append(current_offset)
        current_offset += len(line) + 1  # newline character

    for match in pattern.finditer(content):
        start, end = match.span()

        # Find line number
        line_num = 1
        for idx, offset in enumerate(line_offsets):
            if offset <= start:
                line_num = idx + 1
            else:
                break

        # Generate surrounding context (window of ~60 chars before and after)
        context_start = max(0, start - 60)
        context_end = min(len(content), end + 60)

        prefix = "..." if context_start > 0 else ""
        suffix = "..." if context_end < len(content) else ""

        snippet = prefix + content[context_start:context_end].replace("\n", " ").strip() + suffix

        matches.append({
            "line_number": line_num,
            "matched_text": match.group(0),
            "snippet": snippet,
        })

    return {
        "status": "success",
        "filepath": filepath,
        "keyword": keyword,
        "total_matches": len(matches),
        "matches": matches,
    }


# Standard JSON Schema tool definitions for LLM Tool Calling (OpenAI / Anthropic / Gemini compatible)
TOOLS_SCHEMA = [
    {
        "type": "function",
        "function": {
            "name": "list_files",
            "description": "List files in a given directory, with optional extension filter (e.g. '.pdf', '.docx', '.txt'). Returns file names, paths, sizes, and modified dates.",
            "parameters": {
                "type": "object",
                "properties": {
                    "directory": {
                        "type": "string",
                        "description": "Path to the directory to scan, e.g. './resumes' or an absolute path.",
                    },
                    "extension": {
                        "type": "string",
                        "description": "Optional file extension to filter by (e.g., '.pdf', '.txt', '.docx').",
                    },
                },
                "required": ["directory"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_file",
            "description": "Read and extract all textual content from a PDF, DOCX, or TXT file, returning full content and document metadata.",
            "parameters": {
                "type": "object",
                "properties": {
                    "filepath": {
                        "type": "string",
                        "description": "Relative or absolute path to the file to be read.",
                    },
                },
                "required": ["filepath"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "write_file",
            "description": "Write text content to a file at the specified path. Automatically creates any missing parent directories.",
            "parameters": {
                "type": "object",
                "properties": {
                    "filepath": {
                        "type": "string",
                        "description": "Destination file path (e.g. './summaries/resume_summary.txt').",
                    },
                    "content": {
                        "type": "string",
                        "description": "The textual content to write into the destination file.",
                    },
                },
                "required": ["filepath", "content"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "search_in_file",
            "description": "Perform a case-insensitive search for a keyword or skill inside a document (PDF, DOCX, TXT) and retrieve all matching occurrences with surrounding context snippets.",
            "parameters": {
                "type": "object",
                "properties": {
                    "filepath": {
                        "type": "string",
                        "description": "Path to the file to search within.",
                    },
                    "keyword": {
                        "type": "string",
                        "description": "Keyword, skill, or phrase to locate (e.g. 'Python', 'AWS', 'Docker').",
                    },
                },
                "required": ["filepath", "keyword"],
            },
        },
    },
]

# Function dispatch dictionary
TOOL_MAPPING = {
    "list_files": list_files,
    "read_file": read_file,
    "write_file": write_file,
    "search_in_file": search_in_file,
}
